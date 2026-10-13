import sys
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

IMG = sys.argv[1]; OUT = sys.argv[2]  # e.g. python3 build_slide.py source .
logo = Image.open(f"{IMG}/hennessey-digital-logo.webp").convert("RGB")
# sample the cyan from the D stroke
cy = max((logo.getpixel((x, 900)) for x in range(1100, 1300)), key=lambda p: p[2]-p[0])
BG = logo.getpixel((40, 40)); print("bg", BG, "cyan", cy)
logo.save(f"{OUT}/logo.png")

def circle(src, dst, box):
    im = Image.open(src).convert("RGB").crop(box).resize((900, 900), Image.LANCZOS)
    S = 900; ring = 24
    out = Image.new("RGBA", (S, S), (0, 0, 0, 0))
    m = Image.new("L", (S*4, S*4), 0); ImageDraw.Draw(m).ellipse((0, 0, S*4-1, S*4-1), fill=255)
    m = m.resize((S, S), Image.LANCZOS)
    ringimg = Image.new("RGBA", (S, S), cy + (255,)); out.paste(ringimg, (0, 0), m)
    inner = S - 2*ring
    im = im.resize((inner, inner), Image.LANCZOS)
    m2 = m.resize((inner, inner), Image.LANCZOS)
    out.paste(im, (ring, ring), m2); out.save(dst)

# Dan: 600x338, face centred ~x300 -> square crop
circle(f"{IMG}/dan-hinckley.png", f"{OUT}/dan.png", (131, 0, 469, 338))
# Rob: 1428x2000, face ~x680,y800
circle(f"{IMG}/rob-coleman.webp", f"{OUT}/rob.png", (180, 330, 1230, 1380))

C = lambda t: RGBColor(*t)
WHITE, CYAN, GREY = C((255,255,255)), C(cy), C((190,190,195))
prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
s = prs.slides.add_slide(prs.slide_layouts[6])
s.background.fill.solid(); s.background.fill.fore_color.rgb = C(BG)

def text(x, y, w, h, runs, size, font="Georgia", align=PP_ALIGN.LEFT, spacing=None, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for a in ("margin_left","margin_right","margin_top","margin_bottom"): setattr(tf, a, 0)
    lines = runs if isinstance(runs[0], list) else [runs]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align
        if spacing: p.line_spacing = spacing
        for t, col, bold in line:
            r = p.add_run(); r.text = t; f = r.font
            f.size, f.name, f.bold, f.color.rgb = Pt(size), font, bold, col
    return tb

s.shapes.add_picture(f"{OUT}/logo.png", Inches(0.6), Inches(0.45), Inches(1.35), Inches(1.35))
# cyan accent bar
bar = s.shapes.add_shape(1, Inches(0.75), Inches(2.25), Inches(0.08), Inches(0.42))
bar.fill.solid(); bar.fill.fore_color.rgb = CYAN; bar.line.fill.background()
text(1.0, 2.27, 5, 0.4, [("LIVE WEBINAR", CYAN, True)], 16, "Arial")
text(0.75, 2.85, 6.4, 2.4, [[("AI and the", WHITE, False)], [("Agency of the", WHITE, False)], [("Future", CYAN, False)]], 52, spacing=0.95)
text(0.75, 5.55, 6.4, 0.45, [("Thursday, November 5", WHITE, True)], 24, "Arial")
text(0.75, 6.05, 6.4, 0.45, [("12:00 PM EST", CYAN, True), ("   |   9:00 AM PST", GREY, False)], 20, "Arial")

for name, sub, img, x in [("Dan Hinckley", "Herringbone", "dan.png", 7.35),
                          ("Rob Coleman", "Hennessey Digital", "rob.png", 10.25)]:
    s.shapes.add_picture(f"{OUT}/{img}", Inches(x), Inches(2.0), Inches(2.6), Inches(2.6))
    text(x-0.3, 4.85, 3.2, 0.5, [(name, WHITE, True)], 24, "Arial", PP_ALIGN.CENTER)
    text(x-0.3, 5.35, 3.2, 0.4, [(sub, GREY, False)], 16, "Arial", PP_ALIGN.CENTER)
text(7.05, 1.3, 5.8, 0.4, [("FEATURING", GREY, True)], 13, "Arial", PP_ALIGN.CENTER)
prs.save(f"{OUT}/november_webinar_slide.pptx")
